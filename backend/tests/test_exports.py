import io
import uuid
import pytest
import openpyxl
from bson import ObjectId
from datetime import datetime, timezone, timedelta
from fastapi.testclient import TestClient
from app.main import app
from app.services.excel_service import sanitize_cell_value, generate_applicants_excel

@pytest.fixture
def client():
    with TestClient(app) as tc:
        yield tc

@pytest.fixture
def tpo_token(client):
    res = client.post("/api/auth/login", json={
        "email": "admin@tpo.college.edu",
        "password": "Admin@12345"
    })
    assert res.status_code == 200
    return res.json()["access_token"]

@pytest.fixture
def student_auth(client):
    uid = uuid.uuid4().hex[:6]
    data = {
        "email": f"export_student_{uid}@college.edu",
        "password": "Password123!",
        "full_name": f"Export Student {uid}",
        "roll_number": f"EXP{uid.upper()}",
        "department": "CSE",
        "graduation_year": 2026,
    }
    reg = client.post("/api/auth/register", json=data)
    assert reg.status_code == 201
    token = reg.json()["access_token"]

    client.put(
        "/api/students/profile",
        json={
            "cgpa": 8.75,
            "tenth_percentage": 92.5,
            "twelfth_percentage": 90.0,
            "active_backlogs": 0,
            "department": "CSE",
            "graduation_year": 2026,
            "skills": ["Python", "FastAPI", "React"],
        },
        headers={"Authorization": f"Bearer {token}"},
    )
    return {"token": token, "email": data["email"]}

def test_sanitize_cell_value_formula_injection():
    """Verify formula injection triggers are safely neutralized with leading apostrophe."""
    assert sanitize_cell_value("=1+1") == "'=1+1"
    assert sanitize_cell_value("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_cell_value("-100") == "'-100"
    assert sanitize_cell_value("@SUM(1,2)") == "'@SUM(1,2)"
    # Safe text remains unchanged
    assert sanitize_cell_value("Jane Doe") == "Jane Doe"
    assert sanitize_cell_value(9.5) == 9.5
    assert sanitize_cell_value(None) is None

def test_generate_applicants_excel_direct():
    """Directly test openpyxl workbook generator with mock job and applicants."""
    mock_job = {
        "company_name": "Google India",
        "title": "Software Engineer",
        "salary_package": "24 LPA",
        "job_type": "Full-Time",
        "work_mode": "Hybrid",
        "min_cgpa": 8.0,
        "graduation_year": 2026,
    }

    mock_applications = [
        {
            "status": "Shortlisted",
            "applied_at": datetime.now(timezone.utc),
            "student_snapshot": {
                "full_name": "=FormulaTest Student",  # Malicious formula name to test defense
                "roll_number": "21CS001",
                "email": "student1@college.edu",
                "department": "CSE",
                "cgpa": 8.75,
                "tenth_percentage": 92.5,
                "twelfth_percentage": 89.0,
                "active_backlogs": 0,
                "skills": ["Python", "Algorithms", "React"],
            },
            "status_history": [
                {"status": "Shortlisted", "notes": "Candidate passed round 1"}
            ],
        }
    ]

    stream = generate_applicants_excel(mock_job, mock_applications)
    assert stream is not None

    wb = openpyxl.load_workbook(stream)
    assert "Applicants Roster" in wb.sheetnames
    ws = wb["Applicants Roster"]

    # Check title row
    assert "Google India" in ws["A1"].value
    # Check headers row (Row 5)
    assert ws["A5"].value == "S.No"
    assert ws["B5"].value == "Student Name"
    assert ws["F5"].value == "CGPA"

    # Check student data row (Row 6)
    assert ws["A6"].value == 1
    # Check that formula injection was neutralized
    assert ws["B6"].value == "'=FormulaTest Student"
    assert ws["C6"].value == "21CS001"
    assert ws["F6"].value == 8.75
    assert ws["I6"].value == 0
    assert ws["K6"].value == "Shortlisted"

    # Check frozen panes
    assert ws.freeze_panes == "A6"

def test_export_unauthorized(client):
    """Anonymous requests must be rejected with 401."""
    random_id = str(ObjectId())
    res = client.get(f"/api/exports/applications/{random_id}/excel")
    assert res.status_code == 401

def test_export_forbidden_student(client, student_auth):
    """Students attempting to export administrative rosters must get 403 Forbidden."""
    random_id = str(ObjectId())
    res = client.get(
        f"/api/exports/applications/{random_id}/excel",
        headers={"Authorization": f"Bearer {student_auth['token']}"},
    )
    assert res.status_code == 403

def test_export_invalid_job_id(client, tpo_token):
    """Invalid ObjectId returns 400 Bad Request."""
    res = client.get(
        "/api/exports/applications/invalid-id/excel",
        headers={"Authorization": f"Bearer {tpo_token}"},
    )
    assert res.status_code == 400

def test_export_job_not_found(client, tpo_token):
    """Non-existent job returns 404 Not Found."""
    random_id = str(ObjectId())
    res = client.get(
        f"/api/exports/applications/{random_id}/excel",
        headers={"Authorization": f"Bearer {tpo_token}"},
    )
    assert res.status_code == 404

def test_export_excel_success(client, tpo_token, student_auth):
    """Full workflow: TPO posts drive, student applies, TPO downloads genuine .xlsx."""
    # 1. Post job
    job_payload = {
        "company_name": "Oracle India",
        "title": "Cloud Infrastructure Engineer",
        "description": "OCI Cloud core systems architecture",
        "job_type": "Full-Time",
        "required_branches": ["CSE", "IT"],
        "min_cgpa": 7.0,
        "graduation_year": 2026,
        "max_active_backlogs": 0,
        "required_skills": ["Linux", "Cloud", "Python"],
        "location": "Bangalore",
        "work_mode": "Hybrid",
        "salary_package": "16 LPA",
        "application_deadline": (datetime.now(timezone.utc) + timedelta(days=10)).isoformat(),
        "status": "published",
    }
    create_res = client.post(
        "/api/jobs",
        json=job_payload,
        headers={"Authorization": f"Bearer {tpo_token}"},
    )
    assert create_res.status_code == 201
    job_id = create_res.json()["id"]

    # 2. Student applies
    apply_res = client.post(
        "/api/applications",
        json={"job_id": job_id},
        headers={"Authorization": f"Bearer {student_auth['token']}"},
    )
    assert apply_res.status_code == 201

    # 3. TPO exports excel
    export_res = client.get(
        f"/api/exports/applications/{job_id}/excel",
        headers={"Authorization": f"Bearer {tpo_token}"},
    )
    assert export_res.status_code == 200
    assert (
        export_res.headers["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "Oracle_India" in export_res.headers["Content-Disposition"]
    assert ".xlsx" in export_res.headers["Content-Disposition"]

    # 4. Parse downloaded binary bytes as genuine workbook
    wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
    ws = wb.active
    assert ws.title == "Applicants Roster"
    assert "Oracle India" in ws["A1"].value
    assert ws["A6"].value == 1  # S.No 1
    assert "EXP" in ws["C6"].value  # Roll number
    assert ws["E6"].value == "CSE"  # Department
    assert ws["F6"].value == 8.75  # CGPA
    assert ws["K6"].value == "Applied"  # Status
